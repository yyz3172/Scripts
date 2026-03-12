import re
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

TEST_NAME="0311_vllm_api_stream_chat_multiturn"

LOG_DIR = "/Users/pxy/WorkSpace/logs/" + TEST_NAME
OUTPUT_DIR = "/Users/pxy/WorkSpace/analysis/" + TEST_NAME
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "prefix_cache_hit_rate.xlsx")

os.makedirs(OUTPUT_DIR, exist_ok=True)

# ── Parse log files ──────────────────────────────────────────────────────────
log_files = sorted(f for f in os.listdir(LOG_DIR) if f.startswith("vllm_") and f.endswith(".log"))

data = {}  # {batch_size: [(timestamp, hit_rate), ...]}

for fname in log_files:
    m = re.search(r"_bs(\d+)", fname)
    if not m:
        continue
    batch_size = int(m.group(1))
    entries = []
    with open(os.path.join(LOG_DIR, fname)) as f:
        for line in f:
            if "Prefix cache hit rate:" not in line:
                continue
            ts_m = re.match(r"(\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2})", line)
            # match "Prefix cache hit rate: X%" but not "External Prefix..."
            rate_m = re.search(r"(?<!External )Prefix cache hit rate: ([0-9.]+)%", line)
            if ts_m and rate_m:
                entries.append((ts_m.group(1), float(rate_m.group(1))))
    data[batch_size] = entries

batch_sizes = sorted(data.keys())

# ── Styles ───────────────────────────────────────────────────────────────────
HDR_DARK  = PatternFill("solid", start_color="1F4E79")
HDR_MID   = PatternFill("solid", start_color="2E75B6")
WHITE_FT  = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FT   = Font(name="Arial", size=9)
BOLD_FT   = Font(name="Arial", bold=True, size=10)
CENTER    = Alignment(horizontal="center", vertical="center")

def hdr(ws, row, col, value, fill, width=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = WHITE_FT
    c.fill = fill
    c.alignment = CENTER
    if width:
        ws.column_dimensions[get_column_letter(col)].width = width

# ── Sheet 1: Raw Data ─────────────────────────────────────────────────────────
wb = Workbook()
ws_raw = wb.active
ws_raw.title = "Raw Data"
ws_raw.row_dimensions[1].height = 20

col = 1
batch_col_map = {}
for bs in batch_sizes:
    hdr(ws_raw, 1, col,   f"BS={bs}  Timestamp",    HDR_DARK, width=22)
    hdr(ws_raw, 1, col+1, f"BS={bs}  Hit Rate (%)", HDR_MID,  width=18)
    batch_col_map[bs] = (col, col+1)
    col += 2

for bs in batch_sizes:
    ts_col, rate_col = batch_col_map[bs]
    for r, (ts, rate) in enumerate(data[bs], start=2):
        c_ts   = ws_raw.cell(row=r, column=ts_col,   value=ts)
        c_rate = ws_raw.cell(row=r, column=rate_col, value=rate)
        c_ts.font = BODY_FT
        c_rate.font = BODY_FT
        c_rate.number_format = "0.0"

# ── Sheet 2: Chart Data (index + one column per batch size) ──────────────────
ws_cd = wb.create_sheet("Chart Data")
ws_cd.row_dimensions[1].height = 20

ws_cd.cell(row=1, column=1, value="Sample #").font = BOLD_FT
ws_cd.column_dimensions["A"].width = 12

for idx, bs in enumerate(batch_sizes, start=2):
    c = ws_cd.cell(row=1, column=idx, value=f"BS={bs}")
    c.font = BOLD_FT
    c.alignment = CENTER
    ws_cd.column_dimensions[get_column_letter(idx)].width = 10

max_rows = max(len(data[bs]) for bs in batch_sizes)

for row_i in range(1, max_rows + 1):
    ws_cd.cell(row=row_i + 1, column=1, value=row_i).font = BODY_FT
    for col_i, bs in enumerate(batch_sizes, start=2):
        entries = data[bs]
        if row_i <= len(entries):
            c = ws_cd.cell(row=row_i + 1, column=col_i, value=entries[row_i - 1][1])
            c.font = BODY_FT
            c.number_format = "0.0"

# ── Chart ─────────────────────────────────────────────────────────────────────
chart = LineChart()
chart.title = "Prefix Cache Hit Rate by Batch Size"
chart.style = 10
chart.y_axis.title = "Hit Rate (%)"
chart.x_axis.title = "Sample Index"
chart.y_axis.numFmt = "0.0"
chart.y_axis.scaling.min = 0
chart.y_axis.scaling.max = 100
chart.height = 16
chart.width = 30

for col_i, bs in enumerate(batch_sizes, start=2):
    n = len(data[bs])
    vals = Reference(ws_cd, min_col=col_i, min_row=1, max_row=n + 1)
    chart.add_data(vals, titles_from_data=True)

cats = Reference(ws_cd, min_col=1, min_row=2, max_row=max_rows + 1)
chart.set_categories(cats)

# ── Chart Sheet ───────────────────────────────────────────────────────────────
ws_chart = wb.create_sheet("Chart")
ws_chart.add_chart(chart, "B2")

# Sheet tab order: Raw Data | Chart Data | Chart
wb.save(OUTPUT_FILE)
print(f"Saved: {OUTPUT_FILE}")
